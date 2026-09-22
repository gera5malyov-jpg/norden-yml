#!/usr/bin/env python3
import argparse
import csv
import io
import json
import os
import re
import statistics
import sys
import tempfile
import time
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from urllib.parse import parse_qs, quote, urlencode, urlparse

import requests
from openpyxl import load_workbook

CATALOG_URL = "https://aletan.ru/upload/aletan_catalog_export.csv"
PRICE_PAGE_URL = "https://docs.360.yandex.ru/view/d/_xpJcObcgUxWnp71fA8UZyPegnqahzm72s0qoIz-cKg6VUZsT1dCXzlhdw"
KIT_API = "https://api.kit.yandex.net"
BRAND = "Алетан"
REPORT_PATH = Path("aletan-kit/last_sync_report.json")
STATE_PATH = Path("aletan-kit/state.json")
MIN_INTERVAL_DAYS = 14

def s(value):
    return str(value or "").strip()

def norm(value):
    return " ".join(s(value).casefold().replace("ё", "е").split())

def clean_code(value):
    text = s(value)
    if not text:
        return ""
    if isinstance(value, float) and value.is_integer():
        text = str(int(value))
    return text.replace("\u00a0", " ").strip()

def money(value):
    if value in (None, ""):
        return None
    text = s(value).replace("\u00a0", "").replace(" ", "").replace(",", ".")
    text = re.sub(r"[^0-9.\-]", "", text)
    try:
        d = Decimal(text)
    except (InvalidOperation, ValueError):
        return None
    return d if d > 0 else None

def q2(value):
    return Decimal(value).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)

def load_state():
    if not STATE_PATH.exists():
        return {"last_success_at": None}
    try:
        return json.loads(STATE_PATH.read_text(encoding="utf-8"))
    except Exception:
        return {"last_success_at": None}

def due_for_live(force=False):
    if force:
        return True, None
    raw = s(load_state().get("last_success_at"))
    if not raw:
        return True, None
    try:
        last = datetime.fromisoformat(raw.replace("Z", "+00:00"))
    except Exception:
        return True, None
    age_days = (datetime.now(timezone.utc) - last.astimezone(timezone.utc)).total_seconds() / 86400
    return age_days >= MIN_INTERVAL_DAYS, age_days

def save_state_success():
    STATE_PATH.parent.mkdir(parents=True, exist_ok=True)
    STATE_PATH.write_text(
        json.dumps({"last_success_at": datetime.now(timezone.utc).isoformat()}, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )

def write_report(report):
    REPORT_PATH.parent.mkdir(parents=True, exist_ok=True)
    REPORT_PATH.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))

def decode_csv(content):
    for enc in ("utf-8-sig", "utf-8", "cp1251"):
        try:
            return content.decode(enc)
        except UnicodeDecodeError:
            pass
    return content.decode("utf-8", errors="replace")

def load_catalog():
    r = requests.get(CATALOG_URL, timeout=120, headers={"User-Agent": "Mozilla/5.0 Aletan-KIT-Sync"})
    r.raise_for_status()
    rows = list(csv.DictReader(io.StringIO(decode_csv(r.content)), delimiter=";"))
    by_code, duplicates = {}, {}
    for row in rows:
        if norm(row.get("vendor")) != norm(BRAND):
            continue
        code = clean_code(row.get("vendorCode") or row.get("id"))
        if not code:
            continue
        item = {
            "vendor_code": code,
            "name": s(row.get("name")),
            "catalog_price": money(row.get("price")),
        }
        if code in by_code:
            duplicates.setdefault(code, [by_code[code]]).append(item)
        else:
            by_code[code] = item
    for code in duplicates:
        by_code.pop(code, None)
    return by_code, duplicates, len(rows)

def click_any(page, locators, timeout=4000):
    for loc in locators:
        try:
            if loc.count() and loc.first.is_visible():
                loc.first.click(timeout=timeout)
                return True
        except Exception:
            pass
    return False

def download_price_xlsx(dest):
    from playwright.sync_api import sync_playwright
    with sync_playwright() as p:
        browser = p.chromium.launch(headless=True)
        context = browser.new_context(accept_downloads=True, locale="ru-RU")
        page = context.new_page()
        page.goto(PRICE_PAGE_URL, wait_until="domcontentloaded", timeout=120000)

        # Сначала закрываем баннер cookies на внешней странице.
        for consent in [
            page.get_by_text("Allow essential cookies", exact=True),
            page.get_by_role("button", name=re.compile("Allow essential cookies", re.I)),
            page.get_by_text(re.compile("Только необходимые", re.I)),
            page.get_by_role("button", name=re.compile("необходим", re.I)),
        ]:
            try:
                if consent.count() and consent.first.is_visible():
                    consent.first.click(timeout=5000, force=True)
                    break
            except Exception:
                pass

        page.wait_for_timeout(5000)

        # Редактор таблицы работает внутри iframe volga.yandex.ru/spreadsheet.
        editor = None
        for _ in range(30):
            for frame in page.frames:
                if "volga.yandex.ru/spreadsheet" in frame.url:
                    editor = frame
                    break
            if editor is not None:
                break
            page.wait_for_timeout(500)
        if editor is None:
            raise RuntimeError("Не найден iframe редактора Яндекс Таблиц")

        # Исходник уже является XLSX. В публичном viewer Яндекса его можно
        # получить через session/dv/download/source, используя token и resource-url
        # из iframe. Запрос делаем тем же browser context, чтобы сохранить cookies.
        try:
            parsed = urlparse(editor.url)
            params = parse_qs(parsed.query)
            session_token = (params.get("token") or [""])[0]
            resource_url = (params.get("resource-url") or [""])[0]
            if session_token and resource_url:
                query = urlencode({"mode": "ATTACHEMENT", "url": resource_url})
                source_url = (
                    f"{parsed.scheme}://{parsed.netloc}/session/dv/download/source/"
                    f"{quote(session_token, safe='')}?{query}"
                )
                response = context.request.get(source_url, timeout=120000)
                body = response.body()
                if response.ok and body[:2] == b"PK":
                    dest.write_bytes(body)
                    browser.close()
                    return
        except Exception:
            pass

        page.wait_for_timeout(5000)

        # В публичном режиме Яндекс Таблиц есть отдельная кнопка Download,
        # которая сразу экспортирует XLSX без меню «Файл».
        direct_download = editor.get_by_test_id("Download")
        if direct_download.count() and direct_download.first.is_visible():
            with page.expect_download(timeout=120000) as di:
                direct_download.first.click(timeout=8000)
            di.value.save_as(str(dest))
        else:
            # Fallback для редакторского интерфейса: Файл → Скачать → Excel.
            if not click_any(
                page,
                [
                    editor.get_by_role("button", name=re.compile(r"^Файл$", re.I)),
                    editor.get_by_role("menuitem", name=re.compile(r"^Файл$", re.I)),
                    editor.get_by_text("Файл", exact=True),
                    editor.get_by_role("button", name=re.compile(r"^File$", re.I)),
                    editor.get_by_text("File", exact=True),
                ],
                timeout=5000,
            ):
                visible = editor.locator("button:visible").all_inner_texts()
                try:
                    body_text = editor.locator("body").inner_text(timeout=3000)[:1500]
                except Exception:
                    body_text = ""
                try:
                    test_ids = editor.locator("[data-testid]").evaluate_all(
                        "(els) => els.slice(0, 80).map(e => e.getAttribute('data-testid'))"
                    )
                except Exception:
                    test_ids = []
                frames = [frame.url for frame in page.frames]
                raise RuntimeError(
                    "Не найдена кнопка Download или меню «Файл». "
                    + "title=" + page.title()
                    + "; url=" + page.url
                    + "; frames=" + " || ".join(frames[:10])
                    + "; testids=" + " | ".join(str(x) for x in test_ids[:80])
                    + "; buttons=" + " | ".join(visible[:30])
                    + "; body=" + body_text.replace("\\n", " | ")
                )

            page.wait_for_timeout(800)
            excel = editor.get_by_text(re.compile(r"Microsoft\\s*Excel.*xlsx", re.I))
            if not (excel.count() and excel.first.is_visible()):
                for loc in [
                    editor.get_by_role("menuitem", name=re.compile("Скачать|Download", re.I)),
                    editor.get_by_text(re.compile(r"^Скачать$", re.I)),
                    editor.get_by_text(re.compile(r"^Download$", re.I)),
                ]:
                    try:
                        if loc.count() and loc.first.is_visible():
                            try:
                                loc.first.hover(timeout=4000)
                            except Exception:
                                loc.first.click(timeout=4000)
                            page.wait_for_timeout(1000)
                            break
                    except Exception:
                        pass

            final = None
            for loc in [
                editor.get_by_test_id("ms_excel"),
                editor.get_by_text(re.compile(r"Microsoft\\s*Excel.*xlsx", re.I)),
                editor.get_by_role("menuitem", name=re.compile(r"Microsoft\\s*Excel", re.I)),
            ]:
                try:
                    if loc.count() and loc.first.is_visible():
                        final = loc.first
                        break
                except Exception:
                    pass
            if final is None:
                raise RuntimeError("Не найден пункт «Microsoft Excel (.xlsx)»")

            with page.expect_download(timeout=120000) as di:
                final.click(timeout=8000)
            di.value.save_as(str(dest))
        browser.close()

    if not dest.exists() or dest.stat().st_size < 1000:
        raise RuntimeError("Яндекс Таблица скачалась некорректно")

PRICE_POSITIVE = (
    ("закуп", 120), ("оптов", 100), ("опт", 90), ("дилер", 85),
    ("цена поставщика", 110), ("входная", 80), ("себесто", 70), ("цена", 35),
)
PRICE_NEGATIVE = ("ррц", "рознич", "рекоменд", "маркет", "до скид", "продаж")

def workbook_price_diagnostics(path, catalog):
    """Краткая диагностика ценовых колонок по всем листам прайса."""
    wb = load_workbook(path, read_only=True, data_only=True)
    catalog_codes = set(catalog)
    result = []
    try:
        for ws in wb.worksheets:
            values = [list(row) for row in ws.iter_rows(values_only=True)]
            max_cols = max((len(r) for r in values), default=0)
            if not values or not max_cols:
                continue

            overlap_by_col = []
            for col in range(max_cols):
                found = set()
                for row in values:
                    if col < len(row):
                        code = clean_code(row[col])
                        if code in catalog_codes:
                            found.add(code)
                overlap_by_col.append((len(found), col))
            overlap, sku_col = max(overlap_by_col, default=(0, 0))
            if overlap < 1:
                result.append({"sheet": ws.title, "overlap": overlap})
                continue

            first_match = next(
                (i for i, row in enumerate(values)
                 if sku_col < len(row) and clean_code(row[sku_col]) in catalog_codes),
                None,
            )
            if first_match is None:
                continue

            # Показываем строки непосредственно перед первыми товарами: там обычно
            # находится многоуровневая шапка прайса.
            header_rows = []
            for i in range(max(0, first_match - 8), first_match):
                row = values[i]
                header_rows.append({
                    "row": i + 1,
                    "values": [s(row[col]) if col < len(row) else "" for col in range(max_cols)],
                })

            columns = []
            for col in range(max_cols):
                numeric = 0
                ratios = []
                samples = []
                for row in values[first_match:]:
                    if sku_col >= len(row) or col >= len(row):
                        continue
                    code = clean_code(row[sku_col])
                    if code not in catalog_codes:
                        continue
                    p = money(row[col])
                    if p is None:
                        continue
                    numeric += 1
                    if len(samples) < 5:
                        samples.append({"code": code, "value": str(q2(p))})
                    retail = catalog[code].get("catalog_price")
                    if retail:
                        ratios.append(float(p / retail))
                if numeric:
                    columns.append({
                        "col": col + 1,
                        "numeric_matches": numeric,
                        "median_to_catalog_price": None if not ratios else round(statistics.median(ratios), 4),
                        "samples": samples,
                    })

            result.append({
                "sheet": ws.title,
                "overlap": overlap,
                "sku_col": sku_col + 1,
                "first_data_row": first_match + 1,
                "header_rows": header_rows,
                "numeric_columns": columns,
            })
    finally:
        wb.close()
    return result

def detect_price_map(path, catalog):
    wb = load_workbook(path, read_only=True, data_only=True)
    catalog_codes = set(catalog)
    best = None

    for ws in wb.worksheets:
        values = [list(row) for row in ws.iter_rows(values_only=True)]
        max_cols = max((len(r) for r in values), default=0)
        if not values or not max_cols:
            continue

        overlap_by_col = []
        for col in range(max_cols):
            found = set()
            for row in values:
                if col < len(row):
                    code = clean_code(row[col])
                    if code in catalog_codes:
                        found.add(code)
            overlap_by_col.append((len(found), col))
        overlap, sku_col = max(overlap_by_col, default=(0, 0))
        if overlap < 3:
            continue

        first_match = next(
            (i for i, row in enumerate(values) if sku_col < len(row) and clean_code(row[sku_col]) in catalog_codes),
            None,
        )
        if first_match is None:
            continue

        header_row, header_score = 0, -1
        for i in range(max(0, first_match - 15), first_match):
            row = values[i]
            score = sum(1 for cell in row if s(cell))
            if sku_col < len(row) and any(w in norm(row[sku_col]) for w in ("артикул", "код", "модель", "sku")):
                score += 20
            if score >= header_score:
                header_score, header_row = score, i

        rowh = values[header_row]
        headers = [s(rowh[c]) if c < len(rowh) else "" for c in range(max_cols)]
        candidates = []
        for col, header in enumerate(headers):
            if col == sku_col:
                continue
            hn = norm(header)
            score = max([pts for word, pts in PRICE_POSITIVE if word in hn] or [0])
            if any(word in hn for word in PRICE_NEGATIVE):
                score -= 100
            if score <= 0:
                continue

            covered = numeric = 0
            ratios = []
            for row in values[first_match:]:
                if sku_col >= len(row) or col >= len(row):
                    continue
                code = clean_code(row[sku_col])
                if code not in catalog_codes:
                    continue
                covered += 1
                p = money(row[col])
                if p is None:
                    continue
                numeric += 1
                retail = catalog[code].get("catalog_price")
                if retail:
                    ratios.append(float(p / retail))
            coverage = numeric / max(1, covered)
            if coverage < 0.25:
                continue
            if ratios:
                med = statistics.median(ratios)
                if 0.15 <= med <= 1.05:
                    score += 15
                elif med > 1.5:
                    score -= 30
            candidates.append((score + int(coverage * 20), col, header))

        if not candidates:
            continue

        _, price_col, price_header = max(candidates)
        prices, conflicts = {}, {}
        for row in values[first_match:]:
            if sku_col >= len(row) or price_col >= len(row):
                continue
            code = clean_code(row[sku_col])
            if code not in catalog_codes:
                continue
            p = money(row[price_col])
            if p is None:
                continue
            if code in prices and prices[code] != p:
                conflicts.setdefault(code, {str(prices[code])}).add(str(p))
            else:
                prices[code] = p
        for code in conflicts:
            prices.pop(code, None)

        current = {
            "sheet": ws.title,
            "sku_col": sku_col,
            "price_col": price_col,
            "header_row": header_row,
            "sku_header": headers[sku_col],
            "price_header": price_header,
            "overlap": overlap,
            "prices": prices,
            "conflicts": {k: sorted(v) for k, v in conflicts.items()},
        }
        if best is None or (len(prices), overlap) > (len(best["prices"]), best["overlap"]):
            best = current

    wb.close()
    if best is None or len(best["prices"]) < 3:
        raise RuntimeError("Не удалось надежно определить артикул и закупочную цену в прайсе")
    return best

class KitClient:
    def __init__(self, token):
        self.session = requests.Session()
        self.token = token
        self.last_request_at = 0.0

    def request(self, method, path, *, params=None, body=None, timeout=90):
        url = path if path.startswith("http") else KIT_API + path
        for attempt in range(12):
            delay = 0.55 - (time.monotonic() - self.last_request_at)
            if delay > 0:
                time.sleep(delay)
            self.last_request_at = time.monotonic()
            headers = {"Authorization": "Bearer " + self.token, "Accept": "application/json"}
            try:
                r = self.session.request(method, url, params=params, json=body, headers=headers, timeout=timeout)
            except requests.RequestException:
                if attempt == 11:
                    raise
                time.sleep(min(10, attempt + 1))
                continue
            if r.status_code == 429 or r.status_code >= 500:
                if attempt == 11:
                    r.raise_for_status()
                time.sleep(float(r.headers.get("Retry-After") or min(10, 2 + attempt)))
                continue
            r.raise_for_status()
            return r.json() if r.content else {}
        raise RuntimeError("KIT request retries exhausted")

    def list_all(self, path, params=None, preferred_key=None):
        rows, page = [], 1
        while True:
            query = dict(params or {})
            query.update({"page": page, "per_page": 100})
            payload = self.request("GET", path, params=query)
            batch = payload.get(preferred_key) if preferred_key else None
            if not isinstance(batch, list):
                batch = []
                for key in ("variants", "characteristics", "items", "results", "data"):
                    if isinstance(payload.get(key), list):
                        batch = payload[key]
                        break
            rows.extend(x for x in batch if isinstance(x, dict))
            total = payload.get("total_count") or payload.get("total")
            if not batch or len(batch) < 100 or (total is not None and len(rows) >= int(total)):
                break
            page += 1
        return rows

    def get_variant(self, variant_id):
        return self.request("GET", f"/v1/variants/{variant_id}")

    def update_prices(self, items):
        for start in range(0, len(items), 5000):
            self.request(
                "POST",
                "/v1/variants/prices/bulk_update",
                body={"items": items[start:start + 5000]},
                timeout=180,
            )

def characteristic_values(variant, char_titles):
    out = []
    for row in variant.get("characteristics") or []:
        cid = s(row.get("characteristic_id"))
        title = norm(row.get("title") or char_titles.get(cid))
        if title not in ("артикул поставщика", "артикул", "код для сайта"):
            continue
        vals = []
        if row.get("value") not in (None, ""):
            vals.append(row.get("value"))
        vals.extend(row.get("values") or [])
        out.extend(clean_code(v) for v in vals if clean_code(v))
    return list(dict.fromkeys(out))

def sku_candidates(sku, source_codes):
    sku = clean_code(sku)
    if not sku:
        return []
    out = []
    if sku in source_codes:
        out.append(sku)

    # Проверяем только реальные хвосты SKU, а не весь справочник поставщика.
    # Это покрывает форматы вроде 337-ABC123, KIT-337-ABC-123 и т.п.
    for sep in ("-", "_", "/", " "):
        if sep not in sku:
            continue
        parts = [p for p in sku.split(sep) if p]
        for start in range(1, len(parts)):
            candidate = sep.join(parts[start:])
            if candidate in source_codes:
                out.append(candidate)
        tail = parts[-1] if parts else ""
        if tail in source_codes:
            out.append(tail)
    return list(dict.fromkeys(out))


def build_kit_index(kit, source_codes):
    chars = kit.list_all("/v1/characteristics", {"status": "ACTIVE"}, "characteristics")
    char_titles = {s(x.get("id")): s(x.get("title")) for x in chars if s(x.get("id"))}
    # В KIT параметр brand игнорируется. name выполняет общий поиск и
    # name=Алетан возвращает товары бренда Алетан (проверено API).
    rows = kit.list_all("/v1/variants", {"name": BRAND}, "variants")
    branded = [
        v for v in rows
        if norm(v.get("brand")) == norm(BRAND)
        and s(v.get("status")).upper() != "ARCHIVED"
    ]

    by_code, ambiguous, samples = {}, {}, []
    detailed_reads = 0

    for row in branded:
        vid = s(row.get("id"))
        if not vid:
            continue

        variant = row
        keys = sku_candidates(variant.get("sku"), source_codes)

        # В существующих карточках KIT внутренний SKU часто вида AF-...,
        # а vendorCode Алетан сохранён в названии (например T804BL).
        # Не перебираем весь прайс: вытаскиваем токены из названия и
        # проверяем их за O(1) по set артикулов поставщика.
        product_name = s(variant.get("name"))
        if product_name:
            name_tokens = re.findall(r"[0-9A-Za-zА-Яа-яЁё][0-9A-Za-zА-Яа-яЁё._/-]*", product_name)
            for token in name_tokens:
                token = clean_code(token)
                if token in source_codes:
                    keys.append(token)
                stripped = token.strip("()[]{}.,;:!?")
                if stripped in source_codes:
                    keys.append(stripped)

        # Если в выдаче списка уже есть характеристики — используем их сразу.
        keys.extend(
            k for k in characteristic_values(variant, char_titles)
            if k in source_codes
        )

        # Детальную карточку читаем только когда SKU/списочная выдача
        # не позволили определить артикул поставщика.
        if not keys and not (variant.get("characteristics") or []):
            try:
                variant = kit.get_variant(vid)
                detailed_reads += 1
                keys.extend(sku_candidates(variant.get("sku"), source_codes))
                keys.extend(
                    k for k in characteristic_values(variant, char_titles)
                    if k in source_codes
                )
            except Exception:
                pass

        keys = list(dict.fromkeys(keys))

        if len(samples) < 20:
            samples.append({
                "variant_id": vid,
                "sku": s(variant.get("sku")),
                "name": s(variant.get("name")),
                "brand": s(variant.get("brand")),
                "candidate_codes": keys,
            })

        for code in keys:
            if code in by_code and s(by_code[code].get("id")) != vid:
                ambiguous.setdefault(code, [by_code[code]]).append(variant)
            else:
                by_code[code] = variant

    for code in ambiguous:
        by_code.pop(code, None)

    return by_code, ambiguous, samples, len(branded), detailed_reads

def current_prices(variant):
    p = variant.get("pricing") or {}
    return money(p.get("price")), money(p.get("manual_discount_price"))

def run(args):
    report = {
        "status": "running",
        "dry_run": bool(args.dry_run),
        "brand_rule": "только точный бренд Алетан",
        "creation_rule": "новые товары не создаются",
        "update_scope": "только цены",
        "purchase_price_storage": "закупочная цена в KIT не записывается",
        "price_rule": {"Цена для покупателя": "закупка × 1.30", "Цена до скидки": "закупка × 1.60"},
        "interval_days": MIN_INTERVAL_DAYS,
        "catalog_url": CATALOG_URL,
        "price_page_url": PRICE_PAGE_URL,
        "errors": [],
        "warnings": [],
    }

    if not args.dry_run:
        due, age = due_for_live(args.force)
        report["days_since_last_success"] = None if age is None else round(age, 2)
        if not due:
            report["status"] = "skipped"
            report["reason"] = "С последнего успешного обновления прошло менее 14 дней"
            write_report(report)
            return 0

    token = s(os.environ.get("YANDEX_KIT_TOKEN"))
    if not token:
        raise RuntimeError("YANDEX_KIT_TOKEN is not configured")

    catalog, catalog_duplicates, raw_count = load_catalog()
    report["catalog_rows_total"] = raw_count
    report["catalog_aletan_unique_vendor_codes"] = len(catalog)
    report["catalog_duplicate_vendor_codes"] = list(catalog_duplicates)[:50]

    with tempfile.TemporaryDirectory(prefix="aletan-price-") as td:
        price_file = Path(td) / "aletan-price.xlsx"
        download_price_xlsx(price_file)
        report["price_diagnostics"] = workbook_price_diagnostics(price_file, catalog)
        detected = detect_price_map(price_file, catalog)

    report.update({
        "price_sheet": detected["sheet"],
        "price_header_row_1based": detected["header_row"] + 1,
        "price_sku_column_1based": detected["sku_col"] + 1,
        "price_price_column_1based": detected["price_col"] + 1,
        "price_sku_header": detected["sku_header"],
        "price_purchase_header": detected["price_header"],
        "price_catalog_overlap": detected["overlap"],
        "purchase_prices_unique": len(detected["prices"]),
        "purchase_price_conflicts": detected["conflicts"],
    })

    kit = KitClient(token)
    index, kit_ambiguous, samples, branded_count, detailed_reads = build_kit_index(kit, set(detected["prices"]))
    unsafe_price_header = any(
        word in norm(detected["price_header"])
        for word in ("реком", "рекоменд", "рознич", "ррц", "продаж")
    )
    report["price_source_unsafe"] = unsafe_price_header
    if unsafe_price_header and not args.dry_run:
        raise RuntimeError(
            "Автоблокировка: выбранная колонка прайса похожа на розничную/рекомендованную, а не закупочную"
        )

    report["kit_aletan_active_variants"] = branded_count
    report["kit_detailed_variant_reads"] = detailed_reads
    report["kit_aletan_samples"] = samples
    report["kit_ambiguous_vendor_codes"] = list(kit_ambiguous)[:50]

    matched = sorted(set(detected["prices"]) & set(index))
    report["matched_vendor_codes"] = len(matched)
    report["unmatched_price_vendor_codes"] = sorted(set(detected["prices"]) - set(index))[:100]
    if not matched:
        raise RuntimeError("Не найдено ни одного надежного совпадения Алетан между прайсом и KIT")

    updates, unchanged, examples = [], 0, []
    for code in matched:
        purchase = detected["prices"][code]
        sale = q2(purchase * Decimal("1.30"))
        old = q2(purchase * Decimal("1.60"))
        variant = index[code]
        vid = s(variant.get("id"))
        if "pricing" not in variant:
            variant = kit.get_variant(vid)
        current_old, current_sale = current_prices(variant)

        if len(examples) < 30:
            examples.append({
                "vendor_code": code,
                "kit_sku": s(variant.get("sku")),
                "purchase": str(q2(purchase)),
                "new_customer_price": str(sale),
                "new_old_price": str(old),
                "current_customer_price": None if current_sale is None else str(q2(current_sale)),
                "current_old_price": None if current_old is None else str(q2(current_old)),
            })

        if current_old == old and current_sale == sale:
            unchanged += 1
        else:
            updates.append({"variant_id": vid, "price": str(old), "manual_discount_price": str(sale)})

    report["price_updates_needed"] = len(updates)
    report["prices_unchanged"] = unchanged
    report["examples"] = examples

    if not args.dry_run and updates:
        kit.update_prices(updates)
        report["price_updates_sent"] = len(updates)
    else:
        report["price_updates_sent"] = 0

    report["status"] = "ok"
    if not args.dry_run:
        save_state_success()
    write_report(report)
    return 0

def main():
    parser = argparse.ArgumentParser(description="Алетан: закупочный прайс → существующие товары Yandex KIT")
    parser.add_argument("--dry-run", action="store_true")
    parser.add_argument("--force", action="store_true")
    args = parser.parse_args()
    try:
        return run(args)
    except Exception as exc:
        report = {
            "status": "error",
            "dry_run": bool(args.dry_run),
            "error": str(exc)[:2000],
            "catalog_url": CATALOG_URL,
            "price_page_url": PRICE_PAGE_URL,
        }
        write_report(report)
        return 1

if __name__ == "__main__":
    raise SystemExit(main())
