import csv
import io
from xml.etree import ElementTree as ET


def load_csv(data: bytes, delimiter=None):
    text = data.decode("utf-8-sig")
    sample = text[:4096]
    dialect = csv.Sniffer().sniff(sample, delimiters=delimiter or ",;\t")
    return list(csv.DictReader(io.StringIO(text), dialect=dialect))


def _clean_xml_text(value):
    return (value or "").replace("\xa0", " ").strip()


def _norden_stock_value(value):
    text = _clean_xml_text(value).replace(" ", "").replace(",", ".")
    if not text:
        return 0.0
    if text.startswith(">"):
        text = text[1:]
    try:
        return max(0.0, float(text))
    except ValueError:
        return 0.0


def _load_norden_nomenclature(root):
    rows = []
    for item in root.findall(".//Номенклатура"):
        row = dict(item.attrib)
        free_stock_total = 0.0
        has_free_stock = False
        for child in item:
            tag = child.tag
            text = _clean_xml_text(child.text)
            if tag == "Цена":
                price_type = _clean_xml_text(child.attrib.get("ВидЦен"))
                key = "Цена_" + price_type if price_type else "Цена"
                row[key] = text
            elif tag in {"ОбщийОстаток", "СвободныйОстаток"}:
                warehouse = _clean_xml_text(child.attrib.get("Склад"))
                key = tag + ("_" + warehouse if warehouse else "")
                row[key] = text
                if tag == "СвободныйОстаток":
                    free_stock_total += _norden_stock_value(text)
                    has_free_stock = True
            elif tag == "Заказано":
                month = _clean_xml_text(child.attrib.get("Месяц"))
                row["Заказано_" + month if month else "Заказано"] = text
            elif tag:
                if tag in row:
                    previous = row[tag]
                    row[tag] = previous + [text] if isinstance(previous, list) else [previous, text]
                else:
                    row[tag] = text
        if has_free_stock:
            row["СвободныйОстаток_Итого"] = free_stock_total
        rows.append(row)
    return rows


def load_xml_yml(data: bytes):
    root = ET.fromstring(data)
    rows = []
    for offer in root.findall(".//offer"):
        row = dict(offer.attrib)
        for child in offer:
            if child.tag == "param":
                key = child.attrib.get("name", "").strip()
                if key:
                    row[key] = (child.text or "").strip()
            elif child.tag == "picture":
                row.setdefault("pictures", []).append((child.text or "").strip())
            else:
                row[child.tag] = (child.text or "").strip()
        rows.append(row)
    if rows:
        return rows
    return _load_norden_nomenclature(root)


def load_xlsx(data: bytes, sheet=None):
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb[sheet] if sheet else wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)
    try:
        headers = [str(x).strip() if x is not None else "" for x in next(rows)]
    except StopIteration:
        return []
    return [dict(zip(headers, row)) for row in rows if any(v is not None for v in row)]


def _clean_pdf_cell(value):
    if value is None:
        return ""
    return " ".join(str(value).replace("\xa0", " ").split())


def _unique_headers(headers):
    out = []
    seen = {}
    for index, raw in enumerate(headers, 1):
        base = _clean_pdf_cell(raw) or "column_%d" % index
        count = seen.get(base, 0) + 1
        seen[base] = count
        out.append(base if count == 1 else "%s_%d" % (base, count))
    return out


def load_pdf(data: bytes, required_headers=None):
    """Extract rows from text-based tabular PDFs.

    The first non-empty row of each detected table is treated as its header.
    Only tables containing every configured required header are accepted. This
    deliberately avoids OCR and heuristic guessing: scanned/non-tabular PDFs
    must use a supplier-specific extractor instead of silently importing bad data.
    """
    try:
        import pdfplumber
    except ImportError as exc:
        raise RuntimeError("PDF support requires pdfplumber") from exc

    required = {_clean_pdf_cell(x) for x in (required_headers or []) if _clean_pdf_cell(x)}
    rows = []
    detected_headers = []

    try:
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            for page in pdf.pages:
                for table in page.extract_tables() or []:
                    cleaned = [
                        [_clean_pdf_cell(cell) for cell in row]
                        for row in (table or [])
                        if row and any(_clean_pdf_cell(cell) for cell in row)
                    ]
                    if len(cleaned) < 2:
                        continue
                    raw_headers = cleaned[0]
                    header_set = {x for x in raw_headers if x}
                    detected_headers.append(raw_headers)
                    if required and not required.issubset(header_set):
                        continue
                    headers = _unique_headers(raw_headers)
                    for values in cleaned[1:]:
                        if values == raw_headers:
                            continue
                        if len(values) < len(headers):
                            values = values + [""] * (len(headers) - len(values))
                        elif len(values) > len(headers):
                            values = values[:len(headers)]
                        if not any(values):
                            continue
                        rows.append(dict(zip(headers, values)))
    except Exception as exc:
        if isinstance(exc, RuntimeError):
            raise
        raise RuntimeError("Unable to parse supplier PDF (%s)" % type(exc).__name__) from exc

    if rows:
        return rows

    if detected_headers and required:
        raise RuntimeError(
            "PDF tables found, but required columns are missing: %s"
            % ", ".join(sorted(required))
        )
    raise RuntimeError(
        "No supported text table found in PDF; scanned or non-tabular PDF requires a supplier-specific extractor"
    )
